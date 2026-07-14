"""export_tool.py — Export results to PDF and Word DOCX"""
import os, time

def export_to_pdf(task: str, result: str, agent_type: str, model: str) -> str:
    try:
        from fpdf import FPDF
        class PDF(FPDF):
            def header(self):
                self.set_font("Helvetica","B",14); self.set_text_color(67,97,238)
                self.cell(0,10,"Smart AI Agent — Result",ln=True,align="C")
                self.set_draw_color(67,97,238); self.line(10,self.get_y(),200,self.get_y()); self.ln(4)
            def footer(self):
                self.set_y(-15); self.set_font("Helvetica","I",8); self.set_text_color(150,150,150)
                self.cell(0,10,f"Page {self.page_no()} | Smart AI Agent",align="C")
        pdf = PDF(); pdf.add_page(); pdf.set_auto_page_break(auto=True,margin=15)
        pdf.set_font("Helvetica","B",10); pdf.set_text_color(100,100,100)
        pdf.cell(0,6,f"Agent: {agent_type} | Model: {model} | {time.strftime('%Y-%m-%d %H:%M')}",ln=True); pdf.ln(4)
        pdf.set_font("Helvetica","B",12); pdf.set_text_color(30,30,30); pdf.cell(0,8,"Task:",ln=True)
        pdf.set_font("Helvetica","",11); pdf.set_fill_color(240,244,255); pdf.multi_cell(0,7,task[:500],fill=True); pdf.ln(4)
        pdf.set_font("Helvetica","B",12); pdf.cell(0,8,"Result:",ln=True)
        pdf.set_font("Helvetica","",10); pdf.set_text_color(50,50,50)
        clean = result.replace("**","").replace("##","").replace("#","")
        pdf.multi_cell(0,6,clean[:8000])
        os.makedirs("outputs",exist_ok=True)
        path = f"outputs/result_{int(time.time())}.pdf"
        pdf.output(path); return path
    except Exception: return None

def export_to_docx(task: str, result: str, agent_type: str, model: str) -> str:
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        doc = Document()
        title = doc.add_heading("Smart AI Agent — Result", 0); title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph(f"Agent: {agent_type} | Model: {model} | {time.strftime('%Y-%m-%d %H:%M')}").italic = True
        doc.add_paragraph(); doc.add_heading("Task", level=2); doc.add_paragraph(task)
        doc.add_paragraph(); doc.add_heading("Result", level=2)
        for line in result.split("\n"):
            if line.startswith("## ") or line.startswith("# "):
                doc.add_heading(line.lstrip("#").strip(), level=3)
            elif line.startswith("- ") or line.startswith("• "):
                doc.add_paragraph(line[2:], style="List Bullet")
            elif line.strip():
                doc.add_paragraph(line.replace("**",""))
        os.makedirs("outputs",exist_ok=True)
        path = f"outputs/result_{int(time.time())}.docx"; doc.save(path); return path
    except Exception: return None
