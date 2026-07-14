"""document_tool.py — Read PDF, DOCX, CSV, Excel, images, code files"""
import os

def extract_pdf(path):
    try:
        import fitz
        doc = fitz.open(path)
        pages = []
        for i, page in enumerate(doc):
            t = page.get_text()
            if t.strip(): pages.append(f"--- Page {i+1} ---\n{t.strip()}")
        doc.close()
        full = "\n\n".join(pages)
        return (full[:8000]+"\n[truncated]") if len(full)>8000 else full or "PDF empty."
    except ImportError: return "ERROR: pip install pymupdf"
    except Exception as e: return f"PDF error: {str(e)}"

def extract_docx(path):
    try:
        from docx import Document
        doc = Document(path)
        paras = [p.text for p in doc.paragraphs if p.text.strip()]
        tables = []
        for table in doc.tables:
            for row in table.rows:
                rt = " | ".join(c.text.strip() for c in row.cells)
                if rt.strip(): tables.append(rt)
        full = "\n".join(paras)
        if tables: full += "\n\n[TABLES]\n" + "\n".join(tables)
        return (full[:8000]+"\n[truncated]") if len(full)>8000 else full or "Document empty."
    except ImportError: return "ERROR: pip install python-docx"
    except Exception as e: return f"DOCX error: {str(e)}"

def extract_csv(path):
    try:
        import pandas as pd
        df = pd.read_csv(path)
        return f"CSV: {os.path.basename(path)}\nRows:{len(df)} Cols:{len(df.columns)}\nColumns:{', '.join(df.columns)}\n\nStats:\n{df.describe().to_string()}\n\nFirst 10 rows:\n{df.head(10).to_string()}"[:8000]
    except Exception as e: return f"CSV error: {str(e)}"

def extract_excel(path):
    try:
        import pandas as pd
        xl = pd.ExcelFile(path)
        result = [f"Excel: {os.path.basename(path)}\nSheets: {', '.join(xl.sheet_names)}\n"]
        for sheet in xl.sheet_names[:3]:
            df = pd.read_excel(path, sheet_name=sheet)
            result.append(f"\n--- {sheet} ---\nRows:{len(df)} Cols:{len(df.columns)}\n{df.head(5).to_string()}")
        return "\n".join(result)[:8000]
    except Exception as e: return f"Excel error: {str(e)}"

def extract_image(path):
    try:
        import ollama
        with open(path,"rb") as f: img = f.read()
        resp = ollama.chat(model="llava", messages=[{
            "role":"user","content":"Describe this image in full detail. Include all text, objects, colors, charts.",
            "images":[img]
        }])
        return f"[IMAGE ANALYSIS]\n{resp['message']['content']}"
    except Exception:
        try:
            from PIL import Image
            img = Image.open(path)
            return f"[IMAGE]\nFile:{os.path.basename(path)}\nSize:{img.size[0]}x{img.size[1]}px\nMode:{img.mode}\n\nNote: ollama pull llava for AI vision"
        except Exception as e: return f"Image error: {str(e)}"

def read_code(path):
    try:
        with open(path,"r",encoding="utf-8",errors="replace") as f: content = f.read()
        ext = os.path.splitext(path)[1].lower()
        lang = {".py":"Python",".js":"JavaScript",".ts":"TypeScript",".java":"Java",
                ".cpp":"C++",".c":"C",".cs":"C#",".html":"HTML",".css":"CSS",
                ".sql":"SQL",".r":"R",".go":"Go",".rs":"Rust",".php":"PHP"}.get(ext,"Code")
        header = f"[{lang}: {os.path.basename(path)}]\n"
        return (header+content[:8000]+"\n[truncated]") if len(content)>8000 else header+content
    except Exception as e: return f"Error: {str(e)}"

def process_uploaded_file(path: str) -> dict:
    if not os.path.exists(path):
        return {"type":"error","content":f"Not found: {path}","filename":"","size_kb":0}
    filename = os.path.basename(path)
    size_kb  = round(os.path.getsize(path)/1024, 1)
    ext      = os.path.splitext(filename)[1].lower()
    if ext==".pdf":
        return {"type":"pdf","content":extract_pdf(path),"filename":filename,"size_kb":size_kb}
    elif ext in (".docx",".doc"):
        return {"type":"docx","content":extract_docx(path),"filename":filename,"size_kb":size_kb}
    elif ext==".csv":
        return {"type":"csv","content":extract_csv(path),"filename":filename,"size_kb":size_kb}
    elif ext in (".xlsx",".xls"):
        return {"type":"excel","content":extract_excel(path),"filename":filename,"size_kb":size_kb}
    elif ext in (".jpg",".jpeg",".png",".gif",".bmp",".webp"):
        return {"type":"image","content":extract_image(path),"filename":filename,"size_kb":size_kb}
    elif ext in (".py",".js",".ts",".java",".cpp",".c",".cs",".html",".css",
                 ".sql",".r",".go",".rs",".php",".rb",".kt",".swift",
                 ".sh",".json",".yaml",".yml",".xml",".md",".txt"):
        return {"type":"code","content":read_code(path),"filename":filename,"size_kb":size_kb}
    else:
        try:
            with open(path,"r",encoding="utf-8",errors="replace") as f:
                return {"type":"text","content":f.read(6000),"filename":filename,"size_kb":size_kb}
        except Exception as e:
            return {"type":"unknown","content":str(e),"filename":filename,"size_kb":size_kb}

FILE_ICONS = {"pdf":"📄","docx":"📝","csv":"📊","excel":"📊","image":"🖼️","code":"💻","text":"📃","unknown":"📁","error":"❌"}
